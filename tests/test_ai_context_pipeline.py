from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

ROOT = Path(__file__).resolve().parents[1]

from services.search_docx_AI import extract_doc_items, _determine_type  # noqa: E402
from services.tender_engine.ai import TenderAIReviewer  # noqa: E402
from services.tender_engine.context import ProductDetector  # noqa: E402
from services.tender_engine.data_source import DataTendersRepository  # noqa: E402
from services.tender_engine.docx_reader import DocxTableReader  # noqa: E402
from services.tender_engine.schema import TenderRow  # noqa: E402


def build_fixture(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Технические требования к трансформаторам тока 110 кВ")
    table = doc.add_table(rows=1, cols=4)
    for i, value in enumerate((
        "№ п/п", "Технические требования к оборудованию (наименование параметра)",
        "Требования (значение параметра)", "Предлагаемое участником конкурса"
    )):
        table.rows[0].cells[i].text = value
    for number, parameter, requirement in [
        ("1.1", "Изготовитель", "*"),
        ("1.2", "Заводской тип (марка)", "*"),
        ("1.5", "Цвет внешней изоляции", "белый"),
    ]:
        cells = table.add_row().cells
        cells[0].text = number
        cells[1].text = parameter
        cells[2].text = requirement
        cells[3].text = ""
    doc.save(path)


class AIContextPipelineTests(unittest.TestCase):
    def test_full_document_and_data_tenders_are_separate_contexts(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "fixture.docx"
            build_fixture(path)
            doc = Document(path)
            items, snapshot = extract_doc_items(doc)
            tr_type, voltage = _determine_type(None, items, snapshot)
            self.assertEqual(str(tr_type.name), "ТРГ-УЭТМ-110")
            self.assertEqual(voltage, "110")

            repo = DataTendersRepository(ROOT / "data_tenders")
            data_context = repo.ai_context("ТРГ-УЭТМ-110", max_chars=12000)
            self.assertIn("ТРГ-УЭТМ-110", data_context)
            self.assertIn("nominal_voltage", data_context)
            self.assertIn("dimensions", data_context)
            self.assertIn("mass", data_context)

    def test_only_unresolved_rows_are_sent_to_ai(self):
        rows = [
            TenderRow("a", 0, 1, "1.1", "Изготовитель", "*", "", 3, field_key="manufacturer"),
            TenderRow("b", 0, 2, "1.5", "Цвет внешней изоляции", "белый", "", 3, field_key="external_color"),
        ]
        reviewer = TenderAIReviewer(enabled=True)
        captured = []

        def fake_open_client():
            class Ctx:
                def __enter__(self):
                    return object()
                def __exit__(self, *args):
                    return False
            return Ctx()

        def fake_ask(client, prompt, schema, **kwargs):
            captured.append(prompt)
            return {"results": [{"id": "b", "value": "белый", "confidence": 0.9, "evidence": "ТЗ", "reason": ""}]}

        context = ProductDetector().detect(rows, ["ТРГ-УЭТМ-110"])
        with patch("services.tender_engine.ai.open_client", fake_open_client), patch("services.tender_engine.ai.ask_json_schema", fake_ask):
            result = reviewer.review(rows, context, "DOCUMENT + DATA_TENDERS")

        self.assertEqual(set(result), {"b"})
        self.assertEqual(reviewer.model_calls, 1)
        self.assertEqual(len(captured), 1)
        self.assertIn("data_tenders/БД", captured[0])
        self.assertIn('"id": "a"', captured[0])
        self.assertIn('"id": "b"', captured[0])

    def test_chunk_count_scales_with_prompt_budget(self):
        reviewer = TenderAIReviewer(enabled=False)
        rows = [
            TenderRow(str(i), 0, i, str(i), f"Поле {i}", "*", "", 3, field_key="text")
            for i in range(100)
        ]
        chunks = reviewer._chunks(rows, "x" * 12000)
        self.assertGreaterEqual(len(chunks), 3)
        self.assertTrue(all(1 <= len(chunk) <= 32 for chunk in chunks))


if __name__ == "__main__":
    unittest.main()
