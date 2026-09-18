from __future__ import annotations
import sys, tempfile
from pathlib import Path
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from models.tr_type import Base
from services.tender_engine.docx_reader import DocxTableReader
from services.tender_engine.pipeline import TenderFillingEngine
BAD=Path('/mnt/data/Заполненный_6da5de63-95fe-4730-9fc1-16683ad9b29c Технические требования ПС Нижний Куранах(1).docx')

def blank_answers(src, out):
    doc=Document(src); snap=DocxTableReader().read(doc)
    for r in snap.rows:
        if r.target_cell_index is not None:
            row=doc.tables[r.table_index].rows[r.row_index]
            # use physical cells just like reader
            seen=set(); cells=[]
            for c in row.cells:
                m=id(c._tc)
                if m not in seen: seen.add(m); cells.append(c)
            if r.target_cell_index < len(cells): cells[r.target_cell_index].text=''
    doc.save(out)

with tempfile.TemporaryDirectory() as td:
    td=Path(td); src=td/'blank.docx'; out=td/'filled.docx'; db=td/'x.db'
    blank_answers(BAD,src)
    e=create_engine(f'sqlite:///{db}'); Base.metadata.create_all(e); S=sessionmaker(bind=e)
    with S() as s:
        engine=TenderFillingEngine(ROOT/'data_tenders',ai_enabled=False,overwrite_existing=True)
        res=engine.fill_docx(src,out,s)
        print(res)
        snap=DocxTableReader().read(Document(out))
        print('unresolved rows:')
        for r in snap.rows:
            if not r.current_value: print(r.number,r.parameter,r.requirement)
