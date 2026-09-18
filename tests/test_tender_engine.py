from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]

from models.tr_type import Base  # noqa: E402
from services.tender_engine.data_source import DataTendersRepository  # noqa: E402
from services.tender_engine.docx_reader import DocxTableReader  # noqa: E402
from services.tender_engine.pipeline import TenderFillingEngine  # noqa: E402
from services.tender_engine.schema import TenderRow  # noqa: E402
from services.tender_engine.writer import DocxWriter  # noqa: E402


def build_fixture(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Технические требования к трансформаторам тока 110 кВ")
    table = doc.add_table(rows=1, cols=4)
    hdr = table.rows[0].cells
    hdr[0].text = "№ п/п"
    hdr[1].text = "Технические требования к оборудованию (наименование параметра)"
    hdr[2].text = "Требования (значение параметра)"
    hdr[3].text = "Предлагаемое участником конкурса"
    rows = [
        ("1.1", "Изготовитель", "*"),
        ("1.2.", "Заводской тип (марка)", "*"),
        ("1.3", "Вид внутренней изоляции", "элегаз"),
        ("1.4", "Тип внешней изоляции", "фарфор"),
        ("1.7", "Номинальное напряжение, кВ", "110"),
        ("1.8", "Наибольшее рабочее напряжение, кВ", "126"),
        ("1.9", "Номинальная частота, Гц", "50"),
        ("3.1", "Номинальное давление, МПа", "*"),
        ("4.2", "Масса трансформатора тока /транспортная, кг", "*/*"),
        ("11.1", "Трансформатор тока в сборе", "Да"),
        ("17.3", "Наличие аттестованных производителем специалистов для осуществления гарантийного и постгарантийного ремонтов.", "Да"),
    ]
    for number, parameter, requirement in rows:
        cells = table.add_row().cells
        cells[0].text = number
        cells[1].text = parameter
        cells[2].text = requirement
        cells[3].text = ""
    doc.save(path)


class TenderEngineTests(unittest.TestCase):
    def test_data_tenders_repository_uses_sqlite_and_index(self):
        repo = DataTendersRepository(ROOT / "data_tenders")
        profile = repo.model_profile("ТРГ-УЭТМ-110")
        self.assertEqual(profile["voltage"], "110")
        self.assertTrue(any(x[0] == "110" for x in profile["parameters"]["nominal_voltage"]))
        self.assertTrue(any("2123" in x[0] for x in profile["parameters"]["dimensions"]))
        self.assertTrue(any("SF6" in x[0] for x in profile["parameters"].get("gas_mass", [])))
        self.assertTrue(repo.lookup("Изготовитель", model="ТРГ-УЭТМ-110", requirement="*"))
        self.assertEqual(
            repo.lookup("Номинальное напряжение, кВ", model="ТРГ-УЭТМ-110", requirement="110")[0].value,
            "110",
        )

    def test_reader_keeps_physical_answer_cell(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "fixture.docx"
            build_fixture(path)
            snapshot = DocxTableReader().read(Document(path))
            self.assertTrue(snapshot.rows)
            self.assertTrue(all(row.target_cell_index == 3 for row in snapshot.rows))
            self.assertEqual(next(r for r in snapshot.rows if r.number == "4.2").field_key, "mass")

    def test_writer_highlights_ai_value_only(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=4)
        row = table.rows[0]
        row.cells[3].text = ""
        ai_row = TenderRow(
            "t0_r0_0", 0, 0, "1.3", "Вид внутренней изоляции", "*", "", 3,
            field_key="internal_insulation", proposed_value="Элегаз", source="AI", mark="**"
        )
        DocxWriter().write(doc, [ai_row])
        self.assertEqual(row.cells[3].text, "Элегаз**")
        rpr = row.cells[3].paragraphs[0].runs[0]._r.rPr
        self.assertIsNotNone(rpr)
        self.assertEqual(rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}highlight").get(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
        ), "yellow")

    def test_pipeline_fills_from_data_tenders_before_ai(self):
        with tempfile.TemporaryDirectory() as td:
            td_path = Path(td)
            doc_path = td_path / "fixture.docx"
            out_path = td_path / "filled.docx"
            db_path = td_path / "transformers.db"
            build_fixture(doc_path)
            shutil.copy2(ROOT / "transformers.db", db_path)

            engine = create_engine(f"sqlite:///{db_path}")
            Base.metadata.create_all(engine)
            with engine.begin() as conn:
                conn.execute(text("DELETE FROM tender_parameters"))
                conn.execute(text("DELETE FROM tenders"))
            Session = sessionmaker(bind=engine)
            session = Session()
            try:
                filling = TenderFillingEngine(ROOT / "data_tenders", ai_enabled=False, overwrite_existing=True)
                result = filling.fill_docx(doc_path, out_path, session)
            finally:
                session.close()

            self.assertGreaterEqual(result["filled"], 9)
            self.assertEqual(result["ai_calls"], 0)
            rows = DocxTableReader().read(Document(out_path)).rows
            values = {r.number.rstrip("."): r.current_value for r in rows}
            self.assertIn("Эльмаш", values["1.1"])
            self.assertEqual(values["1.2"], "ТРГ-УЭТМ-110")
            self.assertEqual(values["1.7"], "110")
            self.assertEqual(values["1.8"], "126")
            self.assertEqual(values["1.9"], "50")
            self.assertEqual(values["3.1"], "0,7")
            self.assertEqual(values["11.1"], "Да")
            self.assertEqual(values["17.3"], "Да")
            self.assertEqual(values["4.2"], "")


if __name__ == "__main__":
    unittest.main()
