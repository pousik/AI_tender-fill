"""Чтение технических чертежей из эталонного РЭ.

Модуль отвечает только за одно: при запросе параметров, связанных с
габаритами/массой/размерами, найти в ``data_tenders`` руководство
``1БП.769.001 РЭ``, выделить рисунки А.1-А.6 и передать КАЖДЫЙ рисунок
в GigaChat Vision. Результат кэшируется, поэтому повторный запуск не
повторяет сетевые запросы без изменения исходного изображения.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import subprocess
import zipfile
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:  # pragma: no cover
    Document = None
    qn = None

from services.gigachat_client import ask_vision_json_schema


_REFERENCE_TITLE_RE = re.compile(r"1\s*БП[.\s]*769[.\s]*001\s+РЭ(?:\b|[-—])", re.IGNORECASE)
_FIGURE_RE = re.compile(r"Рисунок\s+А\.(1|2|3|4|5|6)\s*[–—-]\s*(.+)", re.IGNORECASE)

DRAWING_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "figure_id": {"type": "string"},
        "caption": {"type": "string"},
        "model_mentions": {"type": "array", "items": {"type": "string"}},
        "dimensions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "label": {"type": "string"},
                    "value": {"type": "string"},
                    "unit": {"type": "string"},
                    "kind": {"type": "string"},
                    "view": {"type": "string"},
                    "position": {"type": "string"},
                    "evidence": {"type": "string"},
                    "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                },
                "required": ["label", "value", "unit", "kind", "view", "position", "evidence", "confidence"],
                "additionalProperties": False,
            },
        },
        "masses": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "value": {"type": "string"},
                    "unit": {"type": "string"},
                    "context": {"type": "string"},
                    "evidence": {"type": "string"},
                    "confidence": {"type": "number", "minimum": 0, "maximum": 1},
                },
                "required": ["value", "unit", "context", "evidence", "confidence"],
                "additionalProperties": False,
            },
        },
        "notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["figure_id", "caption", "model_mentions", "dimensions", "masses", "notes"],
    "additionalProperties": False,
}


class ReferenceDrawingReader:
    """Находит и один раз распознаёт чертежи А.1-А.6 из 1БП.769.001 РЭ."""

    def __init__(self, root: str | Path):
        self.root = Path(root)
        self.cache_path = self.root / ".drawing_vision_cache.json"

    def read_dimension_drawings(self, *, client=None, force: bool = False) -> dict[str, Any]:
        source = self._find_reference_file()
        if source is None:
            return {
                "source_file": None,
                "figures": [],
                "status": "reference_file_not_found",
            }

        figures = self._extract_dimension_figures(source)
        if not figures:
            return {
                "source_file": source.name,
                "figures": [],
                "status": "dimension_figures_not_found",
            }

        cache = self._load_cache()
        results: list[dict[str, Any]] = []

        for figure in figures:
            raw = figure["image_bytes"]
            digest = hashlib.sha1(raw).hexdigest()
            cache_key = f"{source.resolve()}::{figure['image']}::{digest}"
            cached = cache.get(cache_key)
            if isinstance(cached, dict) and not force:
                result = dict(cached.get("vision") or {})
                result.setdefault("figure_id", figure["figure_id"])
                result.setdefault("caption", figure["caption"])
            else:
                if client is None:
                    raise RuntimeError("Для анализа чертежей необходим открытый GigaChat client.")
                result = self._analyze_figure(client, figure)
                cache[cache_key] = {
                    "source_file": source.name,
                    "image": figure["image"],
                    "figure_id": figure["figure_id"],
                    "caption": figure["caption"],
                    "vision": result,
                }
                self._save_cache(cache)

            result["source_file"] = source.name
            result["image"] = figure["image"]
            result["figure_id"] = figure["figure_id"]
            result["caption"] = figure["caption"]
            results.append(result)

        print(
            f"[GigaChat][DRAWINGS] 1БП.769.001 РЭ: обработано рисунков={len(results)} "
            f"(А.1-А.6); source={source.name}",
            flush=True,
        )
        return {
            "source_file": source.name,
            "figures": results,
            "status": "ok",
        }

    def _find_reference_file(self) -> Path | None:
        if not self.root.exists():
            return None
        candidates: list[Path] = []
        for path in sorted(self.root.rglob("*.docx"), key=lambda p: p.name.lower()):
            try:
                doc = Document(path)
                prefix_parts = [p.text for p in doc.paragraphs[:80]]
                text = "\n".join(prefix_parts)
                if _REFERENCE_TITLE_RE.search(text) and "Руководство по эксплуатации".lower() in text.lower():
                    candidates.append(path)
            except Exception:
                continue
        return candidates[0] if candidates else None

    @staticmethod
    def _relationship_map(doc) -> dict[str, str]:
        mapping: dict[str, str] = {}
        for rel_id, rel in doc.part.rels.items():
            if getattr(rel, "reltype", "") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image":
                target = str(getattr(rel, "target_ref", ""))
                if "media/" in target:
                    mapping[rel_id] = Path(target).name
        return mapping

    @staticmethod
    def _element_images(element, relmap: dict[str, str]) -> list[str]:
        if qn is None:
            return []
        found: list[str] = []
        for blip in element.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if rid and rid in relmap:
                found.append(relmap[rid])
        for data in element.iter("{urn:schemas-microsoft-com:vml}imagedata"):
            rid = data.get(qn("r:id"))
            if rid and rid in relmap:
                found.append(relmap[rid])
        return list(dict.fromkeys(found))

    def _block_sequence(self, doc, relmap: dict[str, str]) -> list[dict[str, Any]]:
        seq: list[dict[str, Any]] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag not in {"p", "tbl"}:
                continue
            text = re.sub(r"\s+", " ", " ".join(child.itertext())).strip()
            images = self._element_images(child, relmap)
            if text or images:
                seq.append({"text": text, "images": images})
        return seq

    def _extract_dimension_figures(self, path: Path) -> list[dict[str, Any]]:
        doc = Document(path)
        relmap = self._relationship_map(doc)
        blocks = self._block_sequence(doc, relmap)
        with zipfile.ZipFile(path, "r") as zf:
            media = {Path(name).name: name for name in zf.namelist() if name.startswith("word/media/")}
            figure_data: list[dict[str, Any]] = []
            for i, block in enumerate(blocks):
                match = _FIGURE_RE.search(block.get("text", ""))
                if not match:
                    continue
                figure_id = f"A.{match.group(1)}"
                caption = f"Рисунок А.{match.group(1)} – {match.group(2).strip()}"
                image_name = None
                # В этом РЭ изображение рисунка находится перед подписью.
                for j in range(i - 1, max(-1, i - 7), -1):
                    imgs = blocks[j].get("images") or []
                    if imgs:
                        image_name = imgs[-1]
                        break
                if not image_name or image_name not in media:
                    continue
                raw = zf.read(media[image_name])
                png_bytes, mime = self._to_png(raw, Path(image_name).suffix)
                figure_data.append({
                    "figure_id": figure_id,
                    "caption": caption,
                    "image": image_name,
                    "image_bytes": png_bytes,
                    "mime_type": mime,
                })
        # Только уникальные А.1-А.6 в документе.
        unique: dict[str, dict[str, Any]] = {}
        for figure in figure_data:
            unique.setdefault(figure["figure_id"], figure)
        return [unique[key] for key in sorted(unique, key=lambda x: int(x.split(".")[1]))]

    @staticmethod
    def _to_png(raw: bytes, suffix: str) -> tuple[bytes, str]:
        suffix = suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            return raw, {
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
                ".webp": "image/webp",
            }[suffix]
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(raw)).convert("RGB")
            out = io.BytesIO()
            image.save(out, format="PNG", optimize=True)
            return out.getvalue(), "image/png"
        except Exception:
            pass
        if suffix in {".wmf", ".emf"}:
            magick = os.getenv("MAGICK_BINARY", "magick")
            proc = subprocess.run(
                [magick, f"{suffix.lstrip('.') }:-", "png:-"],
                input=raw,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=30,
                check=True,
            )
            return proc.stdout, "image/png"
        raise ValueError(f"Не удалось преобразовать изображение {suffix} в PNG.")

    @staticmethod
    def _prompt(figure_id: str, caption: str) -> str:
        return f"""Это технический чертёж из руководства по эксплуатации трансформаторов тока.

Рисунок: {figure_id}
Подпись: {caption}

Твоя задача — прочитать ИМЕННО ГРАФИКУ этого чертежа. Не ограничивайся OCR-текстом.
Определи все размерные линии, стрелки, числа, единицы и их связь с изображённым видом.
Особенно выдели ГАБАРИТНЫЕ РАЗМЕРЫ изделия: общую высоту, ширину, глубину/вынос,
а также массу, если она указана на рисунке или непосредственно связанной с ним таблице.

Для каждого размера укажи:
- точное число;
- единицу;
- вид (спереди/сбоку/сверху/основание/узел);
- назначение размера;
- где на чертеже он расположен;
- краткое доказательство по размерной линии или подписи.

Не придумывай отсутствующие размеры и не превращай номера позиций деталей в размеры.
Если несколько чисел относятся к разным вариантам или моделям — сохрани их все и укажи контекст.
Сохраняй исходные значения дословно. Верни строго структурированный JSON по заданной схеме."""

    def _analyze_figure(self, client, figure: dict[str, Any]) -> dict[str, Any]:
        result = ask_vision_json_schema(
            client,
            figure["image_bytes"],
            self._prompt(figure["figure_id"], figure["caption"]),
            DRAWING_SCHEMA,
            mime_type=figure["mime_type"],
        )
        if not isinstance(result, dict):
            raise RuntimeError(f"Vision для {figure['figure_id']} вернул неверный JSON.")
        return result

    def _load_cache(self) -> dict[str, Any]:
        try:
            raw = json.loads(self.cache_path.read_text(encoding="utf-8"))
            return raw if isinstance(raw, dict) else {}
        except Exception:
            return {}

    def _save_cache(self, cache: dict[str, Any]) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        self.cache_path.write_text(
            json.dumps(cache, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
