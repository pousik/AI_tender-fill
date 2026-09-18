from __future__ import annotations

import json
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document

from .rules import FIELD_ALIASES
from .schema import SourceCandidate
from .text import clean, norm


@dataclass
class SourceDocument:
    path: Path
    text: str
    tables: list[list[list[str]]]


class DataTendersRepository:
    """Единый локальный источник характеристик из ``data_tenders``.

    Источники объединяются без сетевых запросов и в таком порядке:
    1) структурированная SQLite-БД ``data_tenders/tenders.db``;
    2) подготовленный индекс ``.data_tenders_index.json``;
    3) исходные DOCX при отсутствии индекса.

    Все источники преобразуются в один профиль модели, поэтому resolver не
    знает, откуда физически пришло значение.
    """

    MODEL_RE = re.compile(r"ТРГ\s*[-–]?\s*УЭТМ\s*[®™]?\s*[-–]?\s*(35|110|220|330|500|750)", re.I)

    def __init__(self, root: str | Path):
        self.root = Path(root)
        self._docs: list[SourceDocument] = []
        self._profile_cache: dict[str, dict[str, Any]] = {}
        self._db_profiles: dict[str, dict[str, Any]] = {}
        self._index: dict[str, Any] | None = None
        self.refresh()

    def refresh(self) -> None:
        self._docs = []
        self._profile_cache.clear()
        self._db_profiles = self._load_sqlite_profiles()
        self._index = self._load_index()

        # DOCX читаются только как fallback: индекс значительно быстрее на
        # повторных запусках и не требует повторного разбора многомегабайтных ТУ.
        if self._index is None:
            self._docs = self._load_docx_documents()

    def all_text(self, model: str | None = None, limit: int = 250000) -> str:
        chunks: list[str] = []
        if self._index:
            for model_name, profile in self._index.get("profiles", {}).items():
                if model and norm(model_name) != norm(model):
                    continue
                chunks.append(self._format_index_profile(model_name, profile))
        elif self._docs:
            for doc in self.documents_for_model(model):
                chunks.append(f"Файл: {clean(doc.path.name)}\n{doc.text}")
        return "\n\n".join(chunks)[:limit]

    def ai_context(self, model: str | None, *, max_chars: int = 24000) -> str:
        """Компактный контекст БЗ для одного AI-запроса.

        Берутся только факты текущей модели плюс общие нормативные диапазоны,
        а не весь повторяющийся текст корпуса. Это уменьшает prompt и число
        повторных вызовов, не убирая характеристики нужной модели.
        """
        parts: list[str] = []
        profile = self.model_profile(model)
        if profile:
            parts.append(f"[МОДЕЛЬ] {profile.get('model', model)}")
            parts.append(f"[НАПРЯЖЕНИЕ] {profile.get('voltage', '')}")
            for key, values in sorted(profile.get("parameters", {}).items()):
                uniq = []
                for value, evidence, score in values:
                    if value not in uniq:
                        uniq.append(value)
                if uniq:
                    parts.append(f"{key}: {' || '.join(uniq[:8])}")
        # Общие записи без привязки к конкретной модели полезны для диапазонов,
        # классов точности, вторичного тока и других типовых ограничений.
        generic = self._generic_index_records()
        if generic:
            parts.append("[ОБЩИЕ ХАРАКТЕРИСТИКИ]")
            parts.extend(generic[:120])
        return "\n".join(parts)[:max_chars]

    def documents_for_model(self, model: str | None) -> list[SourceDocument]:
        if not model:
            return list(self._docs)
        return [d for d in self._docs if self._doc_matches_model(d, model)]

    def model_profile(self, model: str | None) -> dict[str, Any]:
        if not model:
            return {}
        key = norm(model)
        if key in self._profile_cache:
            return self._profile_cache[key]

        profile: dict[str, Any] = {
            "model": model,
            "manufacturer": "ООО «Эльмаш (УЭТМ)»",
            "parameters": {},
            "facts": [],
        }
        profile["voltage"] = self._voltage(model)

        # 1. Структурированная БД: точные справочные значения.
        self._merge_profile(profile, self._db_profiles.get(key, {}), source_prefix="DATA_TENDERS_DB")

        # 2. Предварительно построенный индекс DOCX.
        if self._index:
            indexed = self._index.get("profiles", {}).get(self._index_model_name(model), {})
            self._merge_index_profile(profile, indexed, model)

        # 3. Fallback на живой разбор DOCX.
        if not self._index:
            for doc in self.documents_for_model(model):
                self._extract_table_facts(doc, model, profile)
                self._extract_paragraph_facts(doc, model, profile)

        self._derive_common_facts(profile)
        self._profile_cache[key] = profile
        return profile

    def lookup(
        self,
        parameter: str,
        *,
        model: str | None,
        voltage: str | None = None,
        requirement: str = "",
        parent_context: str = "",
    ) -> list[SourceCandidate]:
        key = self._field_key(parameter)
        if not model or not key:
            return []
        profile = self.model_profile(model)
        candidates = profile.get("parameters", {}).get(key, [])
        result: list[SourceCandidate] = []
        for value, evidence, score in candidates:
            selected = self._value_for_requirement(key, value, requirement)
            if selected is None:
                continue
            result.append(SourceCandidate(selected, self._source_from_evidence(evidence), score, evidence, model, key))
        result.sort(key=lambda x: x.score, reverse=True)
        return result[:8]

    def _field_key(self, parameter: str) -> str:
        p = norm(parameter)
        # Порядок важен для общих строк с несколькими упоминаниями.
        if "межвитковая изоляция вторичных обмоток" in p:
            return "interturn_test"
        if "изоляция вторичных обмоток должна выдерживать" in p:
            return "secondary_ac_test"
        if "скорость ветра" in p and "отсутствии гололеда" in p:
            return "wind_no_ice"
        if "скорость ветра" in p and "наличии гололеда" in p:
            return "wind_ice"
        if "толщин" in p and "гололеда" in p:
            return "ice_thickness"
        if "абсолютное давление изолирующего газа" in p:
            return "gas_pressure"
        if "масса изолирующего газа" in p or "масса элегаза" in p:
            return "gas_mass"
        for key, aliases in FIELD_ALIASES.items():
            if any(norm(alias) in p for alias in aliases):
                return key
        return ""

    @classmethod
    def _voltage(cls, model: str) -> str | None:
        m = re.search(r"(?:35|110|220|330|500|750)$", clean(model))
        return m.group(0) if m else None

    def _index_model_name(self, model: str) -> str | None:
        wanted = norm(model)
        for name in (self._index or {}).get("profiles", {}):
            if norm(name) == wanted:
                return name
        return None

    def _load_index(self) -> dict[str, Any] | None:
        path = self.root / ".data_tenders_index.json"
        if not path.exists():
            return None
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except Exception as exc:
            print(f"[DATA_TENDERS] индекс не прочитан: {exc}")
            return None

    def _load_sqlite_profiles(self) -> dict[str, dict[str, Any]]:
        path = self.root / "tenders.db"
        if not path.exists():
            return {}
        try:
            con = sqlite3.connect(path)
            con.row_factory = sqlite3.Row
            tables = {row[0] for row in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            profiles: dict[str, dict[str, Any]] = {}

            tr_types = self._table_map(con, "tr_types", "id", "name") if "tr_types" in tables else {}
            volts = self._table_map(con, "voltage_classes", "id", "name") if "voltage_classes" in tables else {}
            isol_types = self._table_map(con, "isol_types", "id", "name") if "isol_types" in tables else {}
            colors = self._table_map(con, "isol_colors", "id", "name") if "isol_colors" in tables else {}
            climats = self._table_map(con, "climats", "id", "name") if "climats" in tables else {}
            acc_table = "accuracy_classses" if "accuracy_classses" in tables else "accuracy_classes"
            accuracies = self._table_map(con, acc_table, "id", "name") if acc_table in tables else {}

            if "gabarits" in tables:
                for row in con.execute("SELECT name, tr_type_voltage FROM gabarits"):
                    voltage = self._voltage_from_pair(row[1], volts)
                    if not voltage:
                        continue
                    model = f"ТРГ-УЭТМ-{voltage}"
                    self._db_add(profiles, model, "dimensions", row[0], "tenders.db:gabarits", 1.0)

            if "weight" in tables:
                for row in con.execute("SELECT value, tr_type_voltage, isol_type, road FROM weight"):
                    voltage = self._voltage_from_pair(row[1], volts)
                    if not voltage:
                        continue
                    model = f"ТРГ-УЭТМ-{voltage}"
                    self._db_add(profiles, model, "mass", row[0], "tenders.db:weight", 0.99)

            for voltage in volts.values():
                model = f"ТРГ-УЭТМ-{voltage}"
                self._db_add(profiles, model, "nominal_voltage", voltage, "tenders.db:voltage_classes", 1.0)

            # Справочники допустимых значений передаются как кандидаты; точное
            # значение выбирается уже по требованию строки.
            for model_voltage in {self._voltage(m) for m in profiles if self._voltage(m)}:
                model = f"ТРГ-УЭТМ-{model_voltage}"
                for value in isol_types.values():
                    self._db_add(profiles, model, "external_insulation", value, "tenders.db:isol_types", 0.90)
                for value in colors.values():
                    self._db_add(profiles, model, "external_color", value, "tenders.db:isol_colors", 0.85)
                for value in climats.values():
                    self._db_add(profiles, model, "climate", value, "tenders.db:climats", 0.85)
                for value in accuracies.values():
                    self._db_add(profiles, model, "accuracy_class", value, "tenders.db:accuracy_classes", 0.80)

            con.close()
            return {norm(k): v for k, v in profiles.items()}
        except Exception as exc:
            print(f"[DATA_TENDERS] SQLite не прочитан: {exc}")
            return {}

    @staticmethod
    def _table_map(con: sqlite3.Connection, table: str, id_col: str, value_col: str) -> dict[int, str]:
        result = {}
        for row in con.execute(f'SELECT "{id_col}", "{value_col}" FROM "{table}"'):
            result[int(row[0])] = clean(row[1])
        return result

    @staticmethod
    def _voltage_from_pair(raw_pair: Any, volts: dict[int, str]) -> str | None:
        try:
            pair = json.loads(raw_pair) if isinstance(raw_pair, str) else raw_pair
        except Exception:
            return None
        if not isinstance(pair, (list, tuple)) or len(pair) < 2:
            return None
        try:
            voltage = volts.get(int(pair[1]))
        except Exception:
            voltage = None
        return clean(voltage) if voltage else None

    @staticmethod
    def _db_add(store: dict[str, dict[str, Any]], model: str, key: str, value: Any, evidence: str, score: float) -> None:
        model_key = norm(model)
        profile = store.setdefault(model_key, {"model": model, "voltage": DataTendersRepository._voltage(model), "parameters": {}, "facts": []})
        value = clean(value)
        if not value:
            return
        bucket = profile.setdefault("parameters", {}).setdefault(key, [])
        if not any(norm(item[0]) == norm(value) for item in bucket):
            bucket.append((value, evidence, score))

    def _merge_profile(self, target: dict[str, Any], source: dict[str, Any], *, source_prefix: str) -> None:
        for key, values in (source.get("parameters", {}) or {}).items():
            for value, evidence, score in values:
                self._add(target, key, value, f"{source_prefix}: {evidence}", score)

    def _merge_index_profile(self, target: dict[str, Any], source: dict[str, Any], model: str) -> None:
        if not source:
            return
        for parameter, values in (source.get("parameters", {}) or {}).items():
            key = self._field_key(parameter)
            if not key:
                continue
            for value in values or []:
                self._add(target, key, value, f"data_tenders index: {parameter}", 0.96)

        # Индекс хранит generic-параметры отдельно от model-profile. Их можно
        # использовать только как ограничения/варианты, но нельзя переносить
        # всю строку в ответ. Например, «1 или 5» превращается в «5», если
        # именно 5 потребовал текущий тендер.
        wanted = norm(model)
        for record in (self._index or {}).get("records", []) or []:
            rec_model = norm(record.get("model") or "")
            if rec_model and rec_model != wanted:
                continue
            parameter = clean(record.get("param_name"))
            value = clean(record.get("value"))
            key = self._field_key(parameter)
            if key and value:
                self._add(target, key, value, f"data_tenders record: {parameter} ({clean(record.get('source_file'))})", 0.93)

            # Газовые таблицы индексируются как «модель -> набор строк».
            # Давление и масса газа извлекаются отдельно, чтобы цифры никогда
            # не переезжали в соседнюю строку тендера. Если для модели есть
            # несколько допустимых газовых смесей, сохраняем все варианты;
            # окончательный выбор для нетипичного поля выполняет AI по полному
            # контексту документа.
            if rec_model == wanted:
                source_text = norm(parameter)
                source_file = clean(record.get('source_file'))
                if value.replace(',', '.') in {"0.7", "0.70"} and "трг" in source_text and "ухл" in source_text:
                    self._add(target, "gas_pressure", self._normalize_numeric(value),
                              f"data_tenders record: газовая таблица ({source_file})", 0.98)
                gas_match = re.fullmatch(
                    r"SF6\s*[–-]\s*([\d,.]+)\s+(CF4|N2)\s*[–-]\s*([\d,.]+)",
                    value, flags=re.I
                )
                if gas_match:
                    gas2 = gas_match.group(2).upper()
                    gas_value = f"SF6 – {gas_match.group(1)}; {gas2} – {gas_match.group(3)}"
                    self._add(target, "gas_mass", gas_value,
                              f"data_tenders record: масса газовой смеси ({source_file})", 0.98)

        # В старом индексе строка ветра имеет имя с двумя условиями и одно
        # значение «40 15». Из неё сразу строятся две семантически разные цели.
        for record in (self._index or {}).get("records", []) or []:
            rec_model = norm(record.get("model") or "")
            if rec_model and rec_model != wanted:
                continue
            parameter = clean(record.get("param_name"))
            value = clean(record.get("value"))
            p = norm(parameter)
            if "максимальная скорость ветра" in p and "отсутствии гололеда" in p:
                nums = re.findall(r"\d+(?:[,.]\d+)?", value)
                if len(nums) >= 2:
                    self._add(target, "wind_no_ice", nums[0], f"data_tenders record: {parameter}", 0.96)
                    self._add(target, "wind_ice", nums[1], f"data_tenders record: {parameter}", 0.96)

    def _derive_common_facts(self, profile: dict[str, Any]) -> None:
        voltage = str(profile.get("voltage") or "")
        params = profile.setdefault("parameters", {})
        if voltage:
            self._add(profile, "manufacturer", "ООО «Эльмаш (УЭТМ)»", "data_tenders: manufacturer", 0.99)
            self._add(profile, "brand", profile.get("model", ""), "data_tenders: model", 1.0)
        if "50" not in [v for v, _, _ in params.get("frequency", [])]:
            self._add(profile, "frequency", "50", "data_tenders: назначение изделия", 0.75)
        if not params.get("internal_insulation"):
            self._add(profile, "internal_insulation", "Элегаз", "data_tenders: газонаполненный ТТ", 0.94)
        if not params.get("external_insulation"):
            self._add(profile, "external_insulation", "Фарфор", "data_tenders: фарфоровая внешняя изоляция", 0.90)

        # В БД климат хранится отдельным справочником, поэтому из УХЛ1
        # детерминированно выводится категория размещения 1.
        if any(norm(v) == "ухл1" for v, _, _ in params.get("climate", [])):
            self._add(profile, "placement_category", "1", "data_tenders.db: УХЛ1 -> категория 1", 0.94)

        # В ТУ скорость ветра записана одной парой «40 15». Разводим её
        # по двум отдельным полям, чтобы значение не уезжало в соседнюю строку.
        wind_values = params.get("wind_no_ice", [])
        if not wind_values and params.get("wind_ice"):
            nums = re.findall(r"\d+(?:[,.]\d+)?", params["wind_ice"][0][0])
            if len(nums) >= 2:
                self._add(profile, "wind_no_ice", nums[0], "data_tenders: ветер без гололеда", 0.96)
                self._add(profile, "wind_ice", nums[1], "data_tenders: ветер с гололедом", 0.96)

    def _generic_index_records(self) -> list[str]:
        result: list[str] = []
        for record in (self._index or {}).get("records", []) or []:
            if record.get("model"):
                continue
            name = clean(record.get("param_name"))
            value = clean(record.get("value"))
            if not name or not value:
                continue
            key = self._field_key(name)
            if key:
                result.append(f"{key}: {name} = {value} | {clean(record.get('source_file'))}")
        # Удаляем только полные дубли строк.
        return list(dict.fromkeys(result))

    def _load_docx_documents(self) -> list[SourceDocument]:
        docs: list[SourceDocument] = []
        if not self.root.exists():
            return docs
        for path in sorted(self.root.glob("*.docx")):
            if path.name.startswith("~$"):
                continue
            try:
                doc = Document(path)
                paragraphs = [clean(p.text) for p in doc.paragraphs if clean(p.text)]
                tables = [
                    [[clean(c.text) for c in row.cells] for row in table.rows]
                    for table in doc.tables
                ]
                text = " ".join(paragraphs)
                for table in tables:
                    for row in table:
                        text += " " + " | ".join(row)
                docs.append(SourceDocument(path, text, tables))
            except Exception as exc:
                print(f"[DATA_TENDERS] пропуск {path.name}: {exc}")
        return docs

    def _doc_matches_model(self, doc: SourceDocument, model: str) -> bool:
        wanted = norm(model).replace(" ", "")
        return any(norm(f"ТРГ-УЭТМ-{m}").replace(" ", "") == wanted for m in self.MODEL_RE.findall(doc.text))

    def _extract_table_facts(self, doc: SourceDocument, model: str, profile: dict[str, Any]) -> None:
        target = norm(model).replace(" ", "")
        for ti, table in enumerate(doc.tables):
            if len(table) < 2:
                continue
            headers = [norm(x) for x in table[1]]
            model_cols = [i for i, x in enumerate(headers) if target in x.replace(" ", "")]
            if not model_cols:
                continue
            col = model_cols[0]
            for ri, row in enumerate(table[2:], 2):
                label = clean(row[0] if row else "")
                value = clean(row[col] if col < len(row) else "")
                self._add_labeled_fact(profile, label, value, f"{doc.path.name}, таблица {ti}, строка {ri}")

        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table):
                if not row:
                    continue
                joined = " | ".join(row)
                if target in norm(joined).replace(" ", ""):
                    self._parse_model_row(row, model, profile, doc.path.name, ti, ri)

    def _add_labeled_fact(self, profile: dict[str, Any], label: str, value: str, evidence: str) -> None:
        key = self._field_key(label)
        if not key or not value:
            return
        # Композитная строка источника «40 15» относится к двум отдельным полям.
        if key in {"wind_no_ice", "wind_ice"}:
            nums = re.findall(r"\d+(?:[,.]\d+)?", value)
            if len(nums) >= 2:
                self._add(profile, "wind_no_ice", nums[0], evidence, 0.96)
                self._add(profile, "wind_ice", nums[1], evidence, 0.96)
                return
        if key == "impulse_voltage" and "одноминутное" in norm(label):
            self._add(profile, "ac_test_voltage", self._first_token(value), evidence, 0.98)
            return
        self._add(profile, key, value, evidence, 0.96)

    def _parse_model_row(self, row: list[str], model: str, profile: dict[str, Any], name: str, ti: int, ri: int) -> None:
        text = norm(" | ".join(row))
        evidence = f"{name}, таблица {ti}, строка {ri}"
        if "абсолютное давление изолирующего газа" in text:
            nums = re.findall(r"\d+(?:[,.]\d+)?", " ".join(row))
            if nums:
                self._add(profile, "gas_pressure", nums[0].replace(".", ","), evidence, 0.99)
        if "масса изолирующего газа" in text or "смесь элегаза" in text:
            gas = self._gas_mass(row)
            if gas:
                self._add(profile, "gas_mass", gas, evidence, 0.98)

    @staticmethod
    def _gas_mass(row: list[str]) -> str:
        text = " ".join(row)
        pairs = re.findall(r"SF6\s*[–-]\s*([\d,.]+).*?(?:N2|CF4)\s*[–-]\s*([\d,.]+)", text, flags=re.I)
        if pairs:
            first, second = pairs[0]
            gas2 = "N2" if re.search(r"N2\s*[–-]", text, re.I) else "CF4"
            return f"SF6 – {first}; {gas2} – {second}"
        return ""

    def _extract_paragraph_facts(self, doc: SourceDocument, model: str, profile: dict[str, Any]) -> None:
        text = doc.text
        self._add(profile, "manufacturer", "ООО «Эльмаш (УЭТМ)»", f"{doc.path.name}: изготовитель", 1.0)
        if model and norm(model) in norm(text):
            self._add(profile, "brand", model, f"{doc.path.name}: обозначение изделия", 1.0)
        if "50 Гц" in text or "50 Гц" in text:
            self._add(profile, "frequency", "50", f"{doc.path.name}: частота", 0.75)
        if "УХЛ1" in text:
            self._add(profile, "climate", "УХЛ1", f"{doc.path.name}: климатическое исполнение", 0.92)
        self._add(profile, "external_insulation", "Фарфор", f"{doc.path.name}: внешняя изоляция", 0.90)
        self._add(profile, "internal_insulation", "Элегаз", f"{doc.path.name}: внутренняя изоляция", 0.97)

    @staticmethod
    def _first_token(value: str) -> str:
        m = re.search(r"[-+]?\d+(?:[,.]\d+)?", value)
        return m.group(0).replace(".", ",") if m else value

    @staticmethod
    def _add(profile: dict[str, Any], key: str, value: str, evidence: str, score: float) -> None:
        value = clean(value)
        if not value:
            return
        bucket = profile.setdefault("parameters", {}).setdefault(key, [])
        if not any(norm(existing[0]) == norm(value) for existing in bucket):
            bucket.append((value, evidence, score))
        profile.setdefault("facts", []).append({"field": key, "value": value, "evidence": evidence, "score": score})

    @staticmethod
    def _source_from_evidence(evidence: str) -> str:
        if "tenders.db:" in evidence:
            return "DATA_TENDERS_DB"
        return "DATA_TENDERS"

    @staticmethod
    def _value_for_requirement(key: str, value: str, requirement: str) -> str | None:
        req = norm(requirement)
        val = clean(value)
        if not val:
            return None
        if not req or re.fullmatch(r"[*\s/]+", req):
            return val
        if req in {"да", "нет"}:
            return val if norm(val) == req else None

        req_token = re.fullmatch(r"[-+]?\d+(?:[,.]\d+)?", clean(requirement))
        if req_token:
            if norm(val) == req:
                return val
            # Из множества «1 или 5», «0,2S; 0,5S ...» выбираем конкретный
            # требуемый элемент, а не записываем всё множество в ячейку.
            options = re.split(r"\s*(?:;|/|\s+или\s+)\s*", val, flags=re.I)
            if any(norm(opt) == req for opt in options):
                return clean(requirement)

        if norm(val) == req:
            return val
        if key in {"climate", "internal_insulation", "external_insulation"}:
            if req in norm(val) or norm(val) in req:
                return val
        if "не более" in req or "не менее" in req:
            nums_req = re.findall(r"\d+(?:[,.]\d+)?", req)
            nums_val = re.findall(r"\d+(?:[,.]\d+)?", norm(val))
            if nums_req and nums_val:
                try:
                    a, b = float(nums_req[0].replace(",", ".")), float(nums_val[0].replace(",", "."))
                    return val if ("не более" in req and b <= a) or ("не менее" in req and b >= a) else None
                except ValueError:
                    return None
        return None

    @staticmethod
    def _normalize_numeric(value: str) -> str:
        text = clean(value).replace('.', ',')
        m = re.fullmatch(r"[-+]?\d+,([0-9]+)", text)
        if m:
            whole = text.rsplit(',', 1)[0]
            frac = m.group(1).rstrip('0')
            return whole if not frac else f"{whole},{frac}"
        return text

    @staticmethod
    def _format_index_profile(model: str, profile: dict[str, Any]) -> str:
        lines = [f"MODEL: {model} | voltage={profile.get('voltage', '')}"]
        for parameter, values in (profile.get("parameters", {}) or {}).items():
            uniq = list(dict.fromkeys(clean(x) for x in values if clean(x)))
            if uniq:
                lines.append(f"PARAM: {parameter} | VALUES: {' || '.join(uniq[:8])}")
        return "\n".join(lines)


def get_data_tenders_knowledge(root: str | Path | None = None) -> DataTendersRepository:
    base = Path(root) if root else Path(__file__).resolve().parents[2] / "data_tenders"
    return DataTendersRepository(base)
