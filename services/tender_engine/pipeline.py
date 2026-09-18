from __future__ import annotations

import json
import os
import time
from pathlib import Path

from docx import Document

from .ai import TenderAIReviewer
from .context import ProductDetector
from .data_source import DataTendersRepository
from .docx_reader import DocxTableReader, snapshot_context
from .resolver import ResolverConfig, ValueResolver
from .schema import Decision, TenderRow
from .template_store import TemplateStore
from .text import clean
from .validator import TenderValidator
from .writer import DocxWriter


class TenderFillingEngine:
    """Полный конвейер: extract -> identify -> template -> facts -> requirements -> AI -> validate -> write."""

    def __init__(self, data_tenders_root: str | Path | None = None, *, ai_enabled: bool | None = None, overwrite_existing: bool = False):
        root = Path(data_tenders_root) if data_tenders_root else Path(__file__).resolve().parents[2] / "data_tenders"
        self.reader = DocxTableReader()
        self.detector = ProductDetector()
        self.repo = DataTendersRepository(root)
        self.resolver = ValueResolver(self.repo, ResolverConfig(overwrite_existing=overwrite_existing))
        self.templates = TemplateStore()
        self.validator = TenderValidator()
        self.ai = TenderAIReviewer(enabled=ai_enabled)
        self.writer = DocxWriter()

    def fill_docx(self, input_path: str | Path, output_path: str | Path, session, *, specialist_name: str | None = None) -> dict:
        started = time.perf_counter()
        source = Path(input_path)
        target = Path(output_path)
        doc = Document(source)
        snapshot = self.reader.read(doc)
        if not snapshot.rows:
            raise ValueError("В документе не найдены строки тендерной таблицы.")
        context = self.detector.detect(snapshot.rows, snapshot.paragraphs)
        if not context.model:
            raise ValueError("Не удалось определить модель/класс изделия.")

        template, template_map, template_score = self.templates.find(
            session, snapshot.rows, filename=source.name, model=context.model
        )
        print(f"[TENDER] rows={len(snapshot.rows)} model={context.model} voltage={context.voltage} template={getattr(template, 'id', None)}")

        unresolved: list[TenderRow] = []
        for row in snapshot.rows:
            template_value = template_map[row.row_id].proposed_value if row.row_id in template_map else ""
            decision = self.resolver.resolve_initial(row, context, template_value=template_value)
            row.proposed_value = clean(decision.value)
            row.source = decision.source
            row.mark = decision.mark if row.source not in {"DB_TEMPLATE", "EXISTING_ANSWER", "DATA_TENDERS"} else ""
            row.confidence = decision.confidence
            row.reason = decision.reason
            row.evidence = list(decision.evidence)
            if not row.proposed_value:
                unresolved.append(row)

        ai_context = self._build_ai_context(snapshot, context)
        ai_results = self.ai.review(unresolved, context, ai_context)
        for row in unresolved:
            result = ai_results.get(row.row_id, {})
            value = clean(result.get("value", ""))
            if not value:
                continue
            confidence = float(result.get("confidence", 0) or 0)
            ok, reason = self.validator.validate(row, value, snapshot.rows)
            if not ok or confidence < 0.55:
                row.reason = reason or f"AI confidence={confidence:.2f} ниже порога"
                continue
            row.proposed_value = value
            row.source = "AI"
            row.mark = "*" if "*" in row.requirement else "**"
            row.confidence = confidence
            row.reason = clean(result.get("reason", ""))
            evidence = clean(result.get("evidence", ""))
            row.evidence = [evidence] if evidence else []

        for row in snapshot.rows:
            if row.proposed_value:
                ok, reason = self.validator.validate(row, row.proposed_value, snapshot.rows)
                if not ok:
                    print(f"[TENDER][VALIDATION] {row.row_id} отклонено: {reason}")
                    row.proposed_value = ""
                    row.mark = ""
                    row.source = "UNRESOLVED"
                    row.reason = reason

        self.writer.write(doc, snapshot.rows)
        target.parent.mkdir(parents=True, exist_ok=True)
        doc.save(target)
        audit = self._audit(source, target, context, template, template_score, snapshot.rows, started)
        target.with_suffix(target.suffix + ".audit.json").write_text(json.dumps(audit, ensure_ascii=True, indent=2), encoding="utf-8")
        return {
            "output": str(target),
            "rows": len(snapshot.rows),
            "filled": sum(bool(r.proposed_value) for r in snapshot.rows),
            "unresolved": [r.row_id for r in snapshot.rows if not r.proposed_value and not r.current_value],
            "template_id": getattr(template, "id", None),
            "template_score": template_score,
            "ai_calls": self.ai.model_calls,
            "detected_type": context.model,
            "voltage": context.voltage,
        }

    def _build_ai_context(self, snapshot, context) -> str:
        lines = ["[DOCUMENT]"]
        lines.extend(snapshot.paragraphs[:80])
        lines.append("[ROWS]")
        for row in snapshot.rows:
            lines.append(
                " | ".join((
                    row.row_id, row.number, row.parent_context, row.parameter,
                    row.requirement, row.current_value, row.proposed_value, row.source
                ))
            )
        lines.append("[PRODUCT]")
        lines.append(f"model={context.model}; voltage={context.voltage}; manufacturer={context.manufacturer}")
        lines.append("[DATA_TENDERS]")
        lines.append(self.repo.ai_context(context.model, max_chars=12000))
        return "\n".join(lines)

    @staticmethod
    def _audit(source, target, context, template, template_score, rows, started):
        return {
            "input": str(source),
            "output": str(target),
            "model": context.model,
            "voltage": context.voltage,
            "template_id": getattr(template, "id", None),
            "template_score": template_score,
            "elapsed_sec": round(time.perf_counter() - started, 3),
            "rows": [r.to_dict() for r in rows],
            "rules": {
                "existing_answer_preserved": True,
                "template_before_ai": True,
                "row_identity": "table,row,physical_cell,number,field_key,parent_context",
                "ai_only_unresolved": True,
                "ai_mark": "* for star requirement, otherwise **",
                "ai_never_trusted_without_validation": True,
                "trained_values_only_after_engineer_save": True,
            },
        }

    def capture_final_docx(self, path: str | Path, session, specialist_name: str | None = None) -> dict[str, int]:
        source = Path(path)
        doc = Document(source)
        snapshot = self.reader.read(doc)
        if not snapshot.answer_columns:
            return {"saved_rows": 0, "template_created": 0, "template_id": 0}
        context = self.detector.detect(snapshot.rows, snapshot.paragraphs)
        for row in snapshot.rows:
            row.proposed_value = clean(row.current_value)
            row.mark = ""
        return self.templates.save(session, snapshot.rows, filename=source.name, model=context.model, voltage=context.voltage, specialist_name=specialist_name)
