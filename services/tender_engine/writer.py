from __future__ import annotations

from copy import deepcopy

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .text import clean, strip_marks


class DocxWriter:
    """Записывает только в физическую ячейку ответа.

    Значения, полученные AI, дополнительно подсвечиваются жёлтым. Значения
    БД/шаблона/явного требования не маркируются и не подсвечиваются.
    """

    def write(self, doc: Document, rows) -> None:
        for row in rows:
            if row.target_cell_index is None or not row.proposed_value:
                continue
            cell = self._cell(doc, row.table_index, row.row_index, row.target_cell_index)
            if cell is None:
                raise RuntimeError(f"Не найдена целевая ячейка: {row.row_id}")
            self._set_cell(cell, row.proposed_value, row.mark, highlight=row.source == "AI")

    @staticmethod
    def _cell(doc: Document, table_index: int, row_index: int, physical_index: int):
        table = doc.tables[table_index]
        row = table.rows[row_index]
        cells, seen = [], set()
        for cell in row.cells:
            marker = id(cell._tc)
            if marker in seen:
                continue
            seen.add(marker)
            cells.append(cell)
        return cells[physical_index] if 0 <= physical_index < len(cells) else None

    @staticmethod
    def _set_cell(cell, value: str, mark: str, *, highlight: bool = False) -> None:
        text = clean(strip_marks(value))
        if text and mark:
            text += mark
        paragraphs = cell.paragraphs or [cell.add_paragraph()]
        first = paragraphs[0]
        style = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
        for p in paragraphs:
            for run in list(p.runs):
                p._p.remove(run._r)
        for p in paragraphs[1:]:
            p._p.getparent().remove(p._p)
        if not text:
            return
        run = first.add_run(text)
        if style is not None:
            run._r.insert(0, deepcopy(style))
        if highlight:
            DocxWriter._set_yellow_highlight(run)

    @staticmethod
    def _set_yellow_highlight(run) -> None:
        r_pr = run._r.get_or_add_rPr()
        existing = r_pr.find(qn("w:highlight"))
        if existing is None:
            existing = OxmlElement("w:highlight")
            r_pr.append(existing)
        existing.set(qn("w:val"), "yellow")
