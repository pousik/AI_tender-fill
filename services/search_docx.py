import os
import re
from docx import Document
from sqlalchemy.orm import Session
from models.tr_type import TrTypeRule, TrType, VoltageClass, Climat, IsolType, IsolColor

MANUFACTURER_NAME = "ООО «Эльмаш (УЭТМ)»"

def extract_options_from_parentheses(text: str) -> list[str]:
    """
    Извлекает варианты из скобок.
    Пример: 'цвет (белый/коричневый) *' -> ['белый', 'коричневый']
    """
    matches = re.findall(r'\((.*?)\)', text)
    options = []
    for match in matches:
        parts = re.split(r'[/,]| или ', match)
        for part in parts:
            cleaned = part.strip().lower()
            if cleaned:
                options.append(cleaned)
    return options


def extract_doc_requirements(doc: Document) -> dict:
    """
    Извлекает базовые технические параметры из таблиц документа.
    """
    requirements = {
        "voltage": None,
        "climat": None,
        "isol_type": None,
        "isol_color": None,
        "raw_text": []
    }

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 3:
                continue

            param_name = cells[1].text.strip().lower()
            required_val = cells[-2].text.strip().lower()
            combined_line = f"{param_name} {required_val}"

            requirements["raw_text"].append(combined_line)

            # 1. Класс напряжения
            if "номинальное напряжение" in param_name:
                nums = re.findall(r'\d+', required_val)
                if nums:
                    requirements["voltage"] = int(nums[0])

            # 2. Климатическое исполнение
            elif "климатическое исполнение" in param_name:
                requirements["climat"] = required_val

            # 3. Изоляция
            elif "изоляции" in param_name and "цвет" not in param_name and "тип" in param_name:
                options = extract_options_from_parentheses(param_name) + extract_options_from_parentheses(required_val)
                requirements["isol_type"] = options if options else required_val

            # 4. Цвет изоляции
            elif "цвет" in param_name and "изоляции" in param_name:
                options = extract_options_from_parentheses(param_name) + extract_options_from_parentheses(required_val)
                requirements["isol_color"] = options if options else required_val

    return requirements


def match_tr_type_by_db_rules(doc: Document, session: Session) -> tuple[TrType, int]:
    """
    Сравнивает параметры из документа с допустимыми значениями из TrTypeRule в БД.
    """
    doc_reqs = extract_doc_requirements(doc)
    doc_text_block = " ".join(doc_reqs["raw_text"]).lower()

    all_rules = session.query(TrTypeRule).all()
    if not all_rules:
        raise ValueError("Таблица правил (TrTypeRule) пуста в БД.")

    matched_type = None

    for rule in all_rules:
        tr_type = session.query(TrType).filter_by(id=rule.tr_type_id).first()
        if not tr_type:
            continue

        valid_voltages = [v.value for v in session.query(VoltageClass).filter(VoltageClass.id.in_(rule.voltage_classes or [])).all()]
        valid_climats = [c.name.lower() for c in session.query(Climat).filter(Climat.id.in_(rule.climats or [])).all()]
        valid_isol_types = [i.name.lower() for i in session.query(IsolType).filter(IsolType.id.in_(rule.isol_types or [])).all()]

        voltage_ok = True
        if doc_reqs["voltage"] and valid_voltages:
            voltage_ok = str(doc_reqs["voltage"]) in valid_voltages

        climat_ok = True
        if (doc_reqs['climat'][-1] != '1'): doc_reqs['climat'] += '1'
        if doc_reqs["climat"] and valid_climats:
            climat_ok = any(c in str(doc_reqs["climat"]) for c in valid_climats) or doc_reqs["climat"] == "*"

        isol_ok = True
        if doc_reqs["isol_type"] and valid_isol_types:
            if isinstance(doc_reqs["isol_type"], list):
                isol_ok = any(opt in valid_isol_types for opt in doc_reqs["isol_type"])
            else:
                isol_ok = any(i in doc_reqs["isol_type"] for i in valid_isol_types) or doc_reqs["isol_type"] == "*"

        name_in_doc = tr_type.name.lower() in doc_text_block

        if voltage_ok and climat_ok and isol_ok:
            matched_type = tr_type
            if name_in_doc:
                break

    if not matched_type:
        raise ValueError("Не удалось подобрать подходящий тип оборудования из БД по параметрам документа.")

    print(f"[MATCH SUCCESS] Определен тип: {matched_type.name} (ID: {matched_type.id})")
    return matched_type, doc_reqs["voltage"]


def process_docx_requirements(docx_path: str, output_path: str, session: Session):
    """
    Анализирует docx файл с техническими требованиями, сопоставляет параметры
    с БД и заполняет предложенные значения.
    """
    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Файл {docx_path} не найден.")

    doc = Document(docx_path)

    # 1. Сопоставляем документ с БД
    tr_type_obj, detected_voltage = match_tr_type_by_db_rules(doc, session)

    # 2. Подгружаем все справочные списки из БД для найденного оборудования
    rule = session.query(TrTypeRule).filter_by(tr_type_id=tr_type_obj.id).first()

    valid_voltages = [v.value for v in session.query(VoltageClass).filter(VoltageClass.id.in_(rule.voltage_classes or [])).all()]
    valid_climats = [c.name.lower() for c in session.query(Climat).filter(Climat.id.in_(rule.climats or [])).all()]
    valid_isol_types = [i.name.lower() for i in session.query(IsolType).filter(IsolType.id.in_(rule.isol_types or [])).all()]
    valid_isol_colors = [c.name.lower() for c in session.query(IsolColor).filter(IsolColor.id.in_(rule.isol_colors or [])).all()]

    # Собираем общий словарь разрешенных значений из БД
    all_valid_db_values = set(
        [str(v) for v in valid_voltages] + 
        valid_climats + 
        valid_isol_types + 
        valid_isol_colors
    )

    voltage_suffix = f"-{detected_voltage}" if detected_voltage else ""
    full_brand_name = f"{tr_type_obj.name}-УЭТМ{voltage_suffix}"

    # 3. Обходим таблицы и заполняем данные
    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 3:
                continue

            param_name = cells[1].text.strip().lower()
            required_val = cells[-2].text.strip()
            target_cell = cells[-1]

            in_brackets_options = extract_options_from_parentheses(param_name) + extract_options_from_parentheses(required_val)

            # Изготовитель
            if "изготовитель" in param_name:
                target_cell.text = MANUFACTURER_NAME
                continue

            # Марка / Заводской тип
            if "заводской тип" in param_name or "марка" in param_name:
                target_cell.text = full_brand_name
                continue

            # Внутренняя изоляция
            if "внутренней изоляции" in param_name:
                matched_types = [t for t in valid_isol_types if t in in_brackets_options or t in required_val.lower()]
                if matched_types:
                    target_cell.text = matched_types[0].capitalize()
                elif valid_isol_types:
                    target_cell.text = valid_isol_types[0].capitalize()
                continue

            # Внешняя изоляция
            if "внешней изоляции" in param_name and "цвет" not in param_name:
                matched_types = [t for t in valid_isol_types if t in in_brackets_options or t in required_val.lower()]
                if matched_types:
                    target_cell.text = matched_types[0].capitalize()
                elif valid_isol_types:
                    target_cell.text = valid_isol_types[0].capitalize()
                continue

            # Цвет изоляции
            if "цвет" in param_name:
                matched_colors = [c for c in valid_isol_colors if c in in_brackets_options or c in required_val.lower()]
                if matched_colors:
                    target_cell.text = ", ".join(c.capitalize() for c in matched_colors)
                elif valid_isol_colors:
                    target_cell.text = ", ".join(c.capitalize() for c in valid_isol_colors)
                continue

            # Напряжение
            if "номинальное напряжение" in param_name:
                nums = re.findall(r'\d+', required_val)
                if nums and nums[0] in valid_voltages:
                    target_cell.text = required_val
                continue

            # Климатическое исполнение
            if "климатическое исполнение" in param_name:
                req_clean = required_val.lower()
                if len(req_clean) > 0:
                    if req_clean[-1] != '1':
                        req_clean += '1'
                if any(cl in req_clean for cl in valid_climats) or required_val == "*":
                    target_cell.text = required_val if required_val != "*" else (valid_climats[0].upper() if valid_climats else required_val)
                continue

            # --- Общая проверка совпадения значений из БД с ТЗ ---
            req_clean = required_val.strip().lower()

            # Если требуемое значение напрямую совпадает с объектом из БД
            if req_clean in all_valid_db_values:
                target_cell.text = required_val
                continue

            # Подстановка для базовых флагов и спецсимволов ("Да", "Нет", "*")
            if required_val.strip() in ["Да", "Нет"] or required_val == "*":
                if required_val != "*":
                    target_cell.text = required_val

    doc.save(output_path)
    print(f"Файл успешно обработан: {output_path} (Марка: {full_brand_name})")