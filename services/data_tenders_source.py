"""Источник справочных констант для базы знаний тендерного модуля.

Раньше допустимые классы напряжения, климатическое исполнение, типы
изоляции и классы точности были жёстко прописаны в Python-коде
(``models/fill_data_db.py``) и дублировали устаревший справочник
``transformers_legacy``. Теперь эти константы извлекаются напрямую из
паспортных docx-документов изделия, лежащих в каталоге ``data_tenders/``
(руководства по эксплуатации / технические условия на трансформаторы).

Если завтра поставщик пришлёт обновлённый паспорт с новым классом
напряжения — достаточно положить новый docx в ``data_tenders/`` и
перезапустить ``bootstrap.py``: код менять не нужно.

Никаких вымышленных значений: если параметр не найден ни в одном
документе каталога, соответствующий список остаётся пустым, а не
подставляется "на всякий случай" из старого справочника.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
DATA_TENDERS_DIR = ROOT / "data_tenders"

_VOLTAGE_RE = re.compile(r"^\d+(?:[.,]\d+)?$")
_CLIMATE_RE = re.compile(r"^[А-ЯЁа-яё]+\d+$")
_ACCURACY_SPLIT_RE = re.compile(r"[;\s]+")
_ACCURACY_MEASURING_RE = re.compile(r"^\d+(?:[.,]\d+)?S?$")
_ACCURACY_PROTECTION_RE = re.compile(r"^\d+PR?$|^TP[YZ]$")
_PAREN_RE = re.compile(r"\(.*?\)")


@dataclass
class ExtractedConstants:
    """Результат разбора всех docx-паспортов каталога ``data_tenders``."""

    voltage_classes: list[str] = field(default_factory=list)
    climats: list[str] = field(default_factory=list)
    isol_types: list[str] = field(default_factory=list)
    isol_colors: list[str] = field(default_factory=list)
    accuracy_classes: list[tuple[str, bool]] = field(default_factory=list)
    sources: list[str] = field(default_factory=list)


def _cell_text(cell) -> str:
    text = cell.text.replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def _dedupe(values: list[str]) -> list[str]:
    """Убирает повторы, которые python-docx даёт для склеенных (merged) ячеек."""
    seen: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.append(value)
    return seen


def _normalize_accuracy_token(token: str) -> str:
    """ГОСТ-обозначения классов точности печатаются латиницей (P), но в
    исходных Word-документах кириллическая «Р» визуально неотличима от
    латинской «P» — нормализуем её, чтобы значения совпадали с тем, что
    участники тендеров пишут в своих ответах."""
    token = token.strip(" ,;.")
    return token.replace("Р", "P").replace("С", "C")


def _extract_from_document(doc, result: ExtractedConstants) -> None:
    # Типы изоляции (материал изолятора) встречаются в тексте описания
    # конструкции и в таблице масс, а не в отдельной построчной таблице —
    # поэтому ищем по всему тексту документа, а не только по заголовкам строк.
    full_text_lower = " ".join(p.text for p in doc.paragraphs).lower()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text_lower += " " + cell.text.lower()

    if "фарфоров" in full_text_lower and "Фарфор" not in result.isol_types:
        result.isol_types.append("Фарфор")
    if "полимерн" in full_text_lower and "Полимер" not in result.isol_types:
        result.isol_types.append("Полимер")

    for table in doc.tables:
        for row in table.rows:
            cells = [_cell_text(c) for c in row.cells]
            if not cells:
                continue
            label, rest = cells[0], _dedupe(cells[1:])
            low = label.lower()

            if low.startswith("номинальное напряжение") and "кв" in low:
                for value in rest:
                    if _VOLTAGE_RE.match(value) and value not in result.voltage_classes:
                        result.voltage_classes.append(value)

            if "климатическое исполнение" in low:
                for value in rest:
                    for token in value.split(","):
                        token = _PAREN_RE.sub("", token).strip()
                        if _CLIMATE_RE.match(token) and token not in result.climats:
                            result.climats.append(token)

            if low.startswith("классы точности"):
                for value in rest:
                    for token in _ACCURACY_SPLIT_RE.split(value):
                        token = _normalize_accuracy_token(token)
                        if not token:
                            continue
                        if _ACCURACY_MEASURING_RE.match(token):
                            existing = {name for name, _ in result.accuracy_classes}
                            if token not in existing:
                                result.accuracy_classes.append((token, True))
                        elif _ACCURACY_PROTECTION_RE.match(token):
                            existing = {name for name, _ in result.accuracy_classes}
                            if token not in existing:
                                result.accuracy_classes.append((token, False))


def extract_constants(data_dir: Path | None = None) -> ExtractedConstants:
    """Разбирает все docx-паспорта каталога ``data_tenders`` и возвращает
    объединённые допустимые значения (объединение по всем найденным файлам,
    без дублей)."""
    data_dir = data_dir or DATA_TENDERS_DIR
    result = ExtractedConstants()
    if not data_dir.exists():
        return result

    for path in sorted(data_dir.glob("*.docx")):
        # Временные файлы Word (~$...) и не-docx пропускаем.
        if path.name.startswith("~$"):
            continue
        try:
            doc = docx.Document(str(path))
        except Exception:
            continue
        _extract_from_document(doc, result)
        result.sources.append(path.name)

    def voltage_key(value: str) -> float:
        try:
            return float(value.replace(",", "."))
        except ValueError:
            return 0.0

    result.voltage_classes.sort(key=voltage_key)
    return result


if __name__ == "__main__":
    data = extract_constants()
    print(f"Источники ({len(data.sources)}): {data.sources}")
    print(f"Классы напряжения: {data.voltage_classes}")
    print(f"Климатическое исполнение: {data.climats}")
    print(f"Типы изоляции: {data.isol_types}")
    print(f"Цвета изоляции: {data.isol_colors}")
    print(f"Классы точности: {data.accuracy_classes}")
